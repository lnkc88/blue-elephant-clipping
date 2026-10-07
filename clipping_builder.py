# -*- coding: utf-8 -*-
"""
BLUE ELEPHANT Daily News Clipping builder

1단계 (후보 추출):
  python clipping_builder.py candidates BlueElephant.xlsx candidates.json
  -> 시트 정리(빈 행/추출 실패 제거, 제목 꼬리표 정리, 중복 병합, 정렬) 후 후보 목록을 출력·저장

2단계 (워드 작성):
  python clipping_builder.py build TEMPLATE.docx candidates.json OUT.docx [--date 2026-10-07] \
      --exclude 3,7,12 --clear-press 5 --merge 31:24
  --date         : 생략하면 한국 시간 기준 오늘 날짜
  --exclude      : 넣지 않을 후보 ID (관련 없는 기사)
  --clear-press  : URL 도메인과 맞지 않아 매체명을 공란으로 만들 ID
  --merge keep:drop : 자동으로 못 합친 같은 기사(원문/네이버)를 수동으로 합침
"""
import sys, re, json, copy, zipfile, shutil, os, datetime, argparse
import openpyxl
from lxml import etree

BE_TABS = ['블루엘리펀트']
COMP_TABS = ['젠틀몬스터', '아이아이컴바인드', '더블러버스', '뭍', 'MUUT', '리끌로우',
             '오클리', '카린', '룩옵티컬', '프로젝트프로덕트']
MKT_TABS = ['아이웨어', '안경', '선글라스', '안경테', '부정경쟁방지법', '지재처 아이웨어']
SECTIONS = [('be', BE_TABS), ('comp', COMP_TABS), ('mkt', MKT_TABS)]
BAD_PRESS = {'새 창 열림'}
FAIL_TITLE = {'', '제목 추출 실패'}

# ---------------- 1단계: 후보 추출 ----------------
SEPS = ['｜', ' | ', ' – ', ' :: ']

def clean_title(t):
    """제목 뒤 매체명·기자명 꼬리표를 떼고 (제목, 꼬리표)를 돌려준다."""
    t = t.strip(); tail = ''
    for sep in SEPS:
        if sep in t:
            head, rest = t.split(sep, 1)
            t = head.strip(); tail = (rest + ' ' + tail).strip()
    m = re.match(r'^(.*\S)\s+-\s+(\S[^-]{0,19})$', t)   # 끝의 " - 매체명" (20자 이하)
    if m and len(m.group(1)) >= 8:
        t = m.group(1).strip(); tail = (m.group(2) + ' ' + tail).strip()
    return t, tail

norm = lambda s: re.sub(r'[^0-9a-zA-Z가-힣]', '', s or '').lower()
is_naver = lambda u: 'naver.com' in u

def parse_date(v):
    if v is None or str(v).strip() == '':
        return None
    if isinstance(v, datetime.datetime):
        return v
    s = str(v).strip()
    m = re.match(r'(\d{4})[-.](\d{1,2})[-.](\d{1,2})', s)
    if not m:
        return None
    y, mo, d = map(int, m.groups())
    hh = mm = 0
    t = re.search(r'(\d{1,2}):(\d{2})', s)
    if t:
        hh, mm = int(t.group(1)), int(t.group(2))
        if '오후' in s and hh < 12: hh += 12
        if '오전' in s and hh == 12: hh = 0
    return datetime.datetime(y, mo, d, hh, mm)

def decide_press(it):
    """매체명 검증. E열 후보를 제목 꼬리표와 대조한다.
    ok: 꼬리표와 일치 / mismatch: 꼬리표와 불일치 -> 공란 / conflict: 중복 행끼리 E열이 서로 다름 -> 공란
    unchecked: 꼬리표가 없어 자동 확인 불가 -> Claude가 URL 도메인으로 확인"""
    cands = [p for p in it['press_cands'] if p]
    tail = norm(it['tail'])
    if tail:
        hit = [p for p in cands if norm(p) and (norm(p) in tail or tail in norm(p))]
        if hit: return hit[0], 'ok'
        return '', ('mismatch' if cands else 'blank')
    uniq = list(dict.fromkeys(cands))
    if len(uniq) == 1: return uniq[0], 'unchecked'
    if len(uniq) > 1: return '', 'conflict'
    return '', 'blank'

def extract(xlsx):
    wb = openpyxl.load_workbook(xlsx)
    items, by_url, no_title = [], {}, []
    for sec, tabs in SECTIONS:
        for ti, tab in enumerate(tabs):
            if tab not in wb.sheetnames:
                continue
            for rno, r in enumerate(wb[tab].iter_rows(min_row=2, max_col=8, values_only=True), start=2):
                r = list(r) + [None] * 8
                date_v, title, press, url = r[0], str(r[1] or '').strip(), str(r[4] or '').strip(), str(r[5] or '').strip()
                if title in FAIL_TITLE:
                    if title == '' and url and url not in [x['url'] for x in no_title]:
                        no_title.append({'tab': tab, 'row': rno, 'url': url})
                    continue
                if not url:
                    continue
                press = '' if press in BAD_PRESS else press
                dt = parse_date(date_v)
                if url in by_url:                      # 같은 URL: 먼저 나온 것 유지, 날짜만 보충, 매체 후보 수집
                    it = by_url[url]
                    if not it['dt'] and dt: it['dt'] = dt
                    if press not in it['press_cands']: it['press_cands'].append(press)
                    if tab != it['tab'] and tab not in it['also']: it['also'].append(tab)
                    continue
                ct, tail = clean_title(title)
                it = dict(section=sec, tab=tab, tab_order=ti, row=rno, dt=dt, title=ct, tail=tail,
                          raw_title=title, press_cands=[press], url=url, also=[], merged_urls=[],
                          content=str(r[6] or '')[:300])
                by_url[url] = it
                items.append(it)
    for it in items:
        it['press'], it['press_check'] = decide_press(it)

    # 같은 기사, 다른 URL: 제목이 같고 매체가 같거나 한쪽이 공란이면 한 행으로 합친다(원문 주소 우선)
    secrank = {'be': 0, 'comp': 1, 'mkt': 2}
    items.sort(key=lambda x: (secrank[x['section']], x['tab_order'], x['row']))
    merged, groups = [], {}
    for it in items:
        k = norm(it['title'])
        g = groups.get(k)
        if g and (not g['press'] or not it['press'] or norm(g['press']) == norm(it['press'])):
            if is_naver(g['url']) and not is_naver(it['url']):
                g['merged_urls'].append(g['url']); g['url'] = it['url']
            else:
                g['merged_urls'].append(it['url'])
            if not g['press'] and it['press']: g['press'], g['press_check'] = it['press'], it['press_check']
            if not g['dt'] and it['dt']: g['dt'] = it['dt']
            for t in [it['tab']] + it['also']:
                if t != g['tab'] and t not in g['also']: g['also'].append(t)
            continue
        if not g: groups[k] = it
        merged.append(it)
    items = merged

    items.sort(key=lambda x: (secrank[x['section']], x['tab_order'],
                              0 if x['dt'] else 1, -(x['dt'].timestamp() if x['dt'] else 0)))
    out = []
    for i, it in enumerate(items, start=1):
        it['id'] = i
        it['date'] = it['dt'].strftime('%y.%m.%d') if it['dt'] else ''
        it['dt'] = it['dt'].isoformat() if it['dt'] else None
        out.append(it)
    return out, no_title

# ---------------- 2단계: 워드 작성 ----------------
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PR = 'http://schemas.openxmlformats.org/package/2006/relationships'
NS = {'w': W}
q = lambda t: '{%s}%s' % (W, t)
WEEK = '월화수목금토일'

def set_cell(tc, text, link_rid=None):
    p = tc.find('w:p', NS)
    for extra in tc.findall('w:p', NS)[1:]:
        tc.remove(extra)
    old_runs = p.findall('.//w:r', NS)
    rpr = copy.deepcopy(old_runs[0].find('w:rPr', NS)) if old_runs else etree.Element(q('rPr'))
    for ch in list(p):
        if ch.tag != q('pPr'):
            p.remove(ch)
    if not text:
        return
    run = etree.Element(q('r'))
    if link_rid:
        st = etree.Element(q('rStyle')); st.set(q('val'), 'ae')
        for old in rpr.findall('w:rStyle', NS): rpr.remove(old)
        rpr.insert(0, st)
    run.append(rpr)
    t = etree.SubElement(run, q('t')); t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    if link_rid:
        h = etree.SubElement(p, q('hyperlink')); h.set('{%s}id' % R, link_rid); h.set(q('history'), '1')
        h.append(run)
    else:
        p.append(run)

def build(template, cand_json, out, date_str, exclude, clear_press=(), merges=()):
    data = json.load(open(cand_json, encoding='utf-8'))
    byid = {x['id']: x for x in data['items']}
    for i in clear_press:
        byid[i]['press'] = ''
    for keep, drop in merges:                          # 수동 병합: drop의 빈칸 정보로 keep을 보충
        k, d = byid[keep], byid[drop]
        if not k['date'] and d['date']: k['date'] = d['date']
        if not k['press'] and d['press'] and drop not in clear_press: k['press'] = d['press']
        exclude = set(exclude) | {drop}
    items = [x for x in data['items'] if x['id'] not in exclude]
    work = out + '_tmp'
    shutil.rmtree(work, ignore_errors=True)
    with zipfile.ZipFile(template) as z: z.extractall(work)
    for root, dirs, files in os.walk(work):
        for f in files:
            if os.path.islink(os.path.join(root, f)): os.remove(os.path.join(root, f))
    dpath, rpath = f'{work}/word/document.xml', f'{work}/word/_rels/document.xml.rels'
    tree = etree.parse(dpath); body = tree.getroot().find('w:body', NS)
    rtree = etree.parse(rpath); rroot = rtree.getroot()
    rid_n = [100]
    def add_link(url):
        rid_n[0] += 1; rid = f'rIdClip{rid_n[0]}'
        e = etree.SubElement(rroot, '{%s}Relationship' % PR)
        e.set('Id', rid); e.set('Type', R + '/hyperlink'); e.set('Target', url); e.set('TargetMode', 'External')
        return rid

    # 날짜: 글자 조각으로 나뉘어 있어도 문단 전체를 합쳐서 찾고 바꾼다
    d = datetime.date.fromisoformat(date_str)
    new_date = f'{d.year}년 {d.month:02d}월 {d.day:02d}일({WEEK[d.weekday()]})'
    pat = re.compile(r'\d{4}\s*년\s*\d{1,2}\s*월\s*\d{1,2}\s*일\s*\(\s*[월화수목금토일]\s*\)')
    replaced = 0
    for p in body.iter(q('p')):
        if p.getparent().tag == q('tc'):          # 표 안 문단은 건너뜀
            continue
        ts = [t for t in p.iter(q('t'))]
        full = ''.join(t.text or '' for t in ts)
        if not pat.search(full):
            continue
        ts[0].text = pat.sub(new_date, full, count=1)
        ts[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        for t in ts[1:]:
            t.text = ''
        replaced += 1
        break
    if not replaced:
        raise SystemExit('오류: 템플릿에서 날짜 문단을 찾지 못했습니다.')

    tbl = body.find('w:tbl', NS)
    rows = tbl.findall('w:tr', NS)

    # 상단 요약 박스: 템플릿 그대로 둔다(수정하지 않음)

    tmpl_be, tmpl_mention, tmpl_tab = rows[2], rows[4], rows[6]
    def na_row(src):
        tr = copy.deepcopy(src); tcs = tr.findall('w:tc', NS)
        for tc in tcs[1:]: tr.remove(tc)
        tcpr = tcs[0].find('w:tcPr', NS)
        tcpr.find('w:tcW', NS).set(q('w'), '10456')
        for old in tcpr.findall('w:gridSpan', NS): tcpr.remove(old)
        gs = etree.Element(q('gridSpan')); gs.set(q('val'), '5'); tcpr.insert(1, gs)
        set_cell(tcs[0], 'N/A'); return tr
    def data_row(src, it, with_tab):
        tr = copy.deepcopy(src); tcs = tr.findall('w:tc', NS)
        vals = [it['date']] + ([it['tab']] if with_tab else []) + [None, it['press'], '온라인']
        for tc, v in zip(tcs, vals):
            if v is None: set_cell(tc, it['title'], add_link(it['url']))
            else: set_cell(tc, v)
        return tr
    plan = [(rows[2], tmpl_be, 'be', False), (rows[4], tmpl_mention, None, False),
            (rows[6], tmpl_tab, 'comp', True), (rows[8], tmpl_tab, 'mkt', True)]
    for placeholder, tmpl, sec, with_tab in plan:
        sec_items = [x for x in items if x['section'] == sec] if sec else []
        new = [data_row(tmpl, x, with_tab) for x in sec_items] or [na_row(tmpl_be)]
        for n in new: placeholder.addprevious(n)
        tbl.remove(placeholder)

    tree.write(dpath, xml_declaration=True, encoding='UTF-8', standalone=True)
    rtree.write(rpath, xml_declaration=True, encoding='UTF-8', standalone=True)
    if os.path.exists(out): os.remove(out)
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for root, dirs, files in os.walk(work):
            for f in files:
                full = os.path.join(root, f); z.write(full, os.path.relpath(full, work))
    shutil.rmtree(work)
    with zipfile.ZipFile(out) as z:
        doc_text = ''.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', z.read('word/document.xml').decode('utf-8')))
    if new_date not in doc_text:
        raise SystemExit(f'오류: 완성 파일에 날짜 "{new_date}"가 없습니다.')
    print(f'날짜 확인: {new_date}')
    return items

if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('mode', choices=['candidates', 'build'])
    ap.add_argument('args', nargs='+')
    ap.add_argument('--date', help='YYYY-MM-DD, 생략하면 한국 시간 오늘'); ap.add_argument('--exclude', default='')
    ap.add_argument('--clear-press', default='', help='매체명을 공란으로 만들 ID (예: 3,7)')
    ap.add_argument('--merge', default='', help='같은 기사 수동 병합 keep:drop (예: 31:24)')
    a = ap.parse_args()
    if a.mode == 'candidates':
        xlsx, out = a.args
        items, no_title = extract(xlsx)
        json.dump({'items': items, 'no_title': no_title}, open(out, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
        for x in items:
            also = f" (+{','.join(x['also'])})" if x['also'] else ''
            chk = {'ok': '', 'unchecked': ' ※매체 확인필요', 'mismatch': ' ※매체 불일치→공란', 'conflict': ' ※매체 충돌→공란', 'blank': ''}[x['press_check']]
            mg = f" [병합: {len(x['merged_urls'])}건]" if x['merged_urls'] else ''
            print(f"[{x['id']}] {x['section']}|{x['tab']}{also}|{x['date']}|{x['title']}|{x['press'] or '(공란)'}{chk}{mg}|{x['url']}")
            print(f"      {x['content'][:150]!r}")
        print('\n제목 없이 URL만 있는 행:'); [print('  ', n) for n in no_title]
    else:
        template, cand, out = a.args
        ids = lambda s: {int(v) for v in s.split(',') if v.strip()}
        mg = [tuple(int(v) for v in p.split(':')) for p in a.merge.split(',') if p.strip()]
        if not a.date:
            a.date = (datetime.datetime.now(datetime.timezone.utc) + datetime.timedelta(hours=9)).date().isoformat()  # 한국 시간 오늘
        used = build(template, cand, out, a.date, ids(a.exclude), ids(a.clear_press), mg)
        print(f'완료: {out} / 기사 {len(used)}건')
