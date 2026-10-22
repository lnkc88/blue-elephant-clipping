# -*- coding: utf-8 -*-
"""
YuhanClorox Weekly News Clipping builder

0단계 (시트 내려받기):
  python yuhan_builder.py download SHEET_ID_또는_URL YuhanClorox.xlsx
  -> 구글 시트(링크 공유 '뷰어')의 모든 기사 탭(1행이 date,title,...,url,...인 탭)을 xlsx 하나로 묶음

1단계 (후보 추출):
  python yuhan_builder.py candidates YuhanClorox.xlsx cand.json
  -> 빈 행/제목 추출 실패 제거, 제목 꼬리표 정리, 중복(같은 URL·같은 제목) 병합 후 후보 목록 출력·저장

2단계 (워드 작성):
  python yuhan_builder.py build TEMPLATE.docx cand.json assign.json OUT.docx [--date 2026-10-26]
  assign.json 예:
  {
    "own":     [3, 5],                                  # 유한클로락스 (순서대로)
    "mention": [8],                                     # 언급기사
    "comp":    [[12, "제품&서비스"], [15, "사회공헌"]],  # 경쟁사 [ID, 구분]
    "mkt":     [20, 22, 25],                            # 시장/업계 동향 (구분은 자동으로 '기타')
                                                        # 기사 0건인 섹션도 커버 행(라벨)은 항상 남김
    "clear_press": [22],                                # 매체명을 공란으로 만들 ID
    "merge": [[20, 31]]                                 # 같은 기사 수동 병합 [남길 ID, 합칠 ID]
  }
  --date 생략 시 한국 시간 오늘
"""
import sys, re, json, copy, zipfile, shutil, os, datetime, argparse
import openpyxl
from lxml import etree

HEADER = ['date', 'title', 'summary', 'name', 'press', 'url', 'content', 'source']
BAD_PRESS = {'새 창 열림'}
FAIL_TITLE = {'', '제목 추출 실패'}
COMP_CATS = ['사업전략', '제품&서비스', '사회공헌', '기타']

# ---------------- 0단계: 시트 내려받기 ----------------
def download(sheet, out_xlsx):
    import csv, io, html, urllib.request
    m = re.search(r'/d/([A-Za-z0-9_-]+)', sheet)
    sid = m.group(1) if m else sheet.strip()
    base = f'https://docs.google.com/spreadsheets/d/{sid}'
    def get(url):
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=120) as r:
            if r.geturl().split('/')[2] != 'docs.google.com':
                raise SystemExit('오류: 시트가 로그인 페이지로 넘어갔습니다. 공유 설정이 "링크가 있는 모든 사용자 - 뷰어"인지 확인하세요.')
            return r.read().decode('utf-8')
    page = get(base + '/htmlview')
    tabs = [(html.unescape(n), g) for n, g in re.findall(r'\{name: "([^"]+)", pageUrl: "[^"]*gid=(\d+)', page)]
    if not tabs:
        raise SystemExit('오류: 탭 목록을 읽지 못했습니다. 공유 설정을 확인하세요.')
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    used, skipped = 0, []
    for t, gid in tabs:
        text = get(f'{base}/gviz/tq?tqx=out:csv&gid={gid}')
        rows = list(csv.reader(io.StringIO(text)))
        if not rows or [c.strip().lower() for c in rows[0][:8]] != HEADER:
            skipped.append(t); continue
        ws = wb.create_sheet(t[:31])
        for r in rows:
            ws.append([v[:32000] for v in (r + [''] * 8)[:8]])
        n = sum(1 for r in rows[1:] if any(v.strip() for v in r[:8]))
        print(f'  {t}: {n}행'); used += 1
    if not used:
        raise SystemExit('오류: 기사 탭을 하나도 찾지 못했습니다.')
    wb.save(out_xlsx)
    print(f'내려받기 완료: {out_xlsx} (기사 탭 {used}개' + (f', 제외한 탭: {skipped}' if skipped else '') + ')')

# ---------------- 1단계: 후보 추출 ----------------
SEPS = ['｜', ' | ', ' – ', ' :: ']

def clean_title(t):
    t = t.strip(); tail = ''
    for sep in SEPS:
        if sep in t:
            head, rest = t.split(sep, 1)
            t = head.strip(); tail = (rest + ' ' + tail).strip()
    m = re.match(r'^(.*\S)\s+-\s+(\S[^-]{0,19})$', t)
    if m and len(m.group(1)) >= 8:
        t = m.group(1).strip(); tail = (m.group(2) + ' ' + tail).strip()
    return t, tail

norm = lambda s: re.sub(r'[^0-9a-zA-Z가-힣]', '', s or '').lower()
is_naver = lambda u: 'naver.com' in u

def parse_date(v):
    if v is None or str(v).strip() == '':
        return None
    if isinstance(v, datetime.datetime):
        return v
    s = str(v).strip()
    m = re.match(r'(\d{4})[-.](\d{1,2})[-.](\d{1,2})', s)
    if not m:
        return None
    y, mo, d = map(int, m.groups())
    hh = mm = 0
    t = re.search(r'(\d{1,2}):(\d{2})', s)
    if t:
        hh, mm = int(t.group(1)), int(t.group(2))
        if '오후' in s and hh < 12: hh += 12
        if '오전' in s and hh == 12: hh = 0
    return datetime.datetime(y, mo, d, hh, mm)

def decide_press(it):
    cands = [p for p in it['press_cands'] if p]
    tail = norm(it['tail'])
    if tail:
        hit = [p for p in cands if norm(p) and (norm(p) in tail or tail in norm(p))]
        if hit: return hit[0], 'ok'
        return '', ('mismatch' if cands else 'blank')
    uniq = list(dict.fromkeys(cands))
    if len(uniq) == 1: return uniq[0], 'unchecked'
    if len(uniq) > 1: return '', 'conflict'
    return '', 'blank'

def extract(xlsx):
    wb = openpyxl.load_workbook(xlsx)
    items, by_url, no_title = [], {}, []
    for ti, tab in enumerate(wb.sheetnames):
        for rno, r in enumerate(wb[tab].iter_rows(min_row=2, max_col=8, values_only=True), start=2):
            r = list(r) + [None] * 8
            date_v, title, press, url = r[0], str(r[1] or '').strip(), str(r[4] or '').strip(), str(r[5] or '').strip()
            if title in FAIL_TITLE or title.replace('\xa0', ' ').strip().startswith('제목 추출 실패'):
                if url and url not in [x['url'] for x in no_title]:   # 제목은 없지만 URL이 있는 행은 보고용으로 모음
                    no_title.append({'tab': tab, 'row': rno, 'url': url})
                continue
            if not url:
                continue
            press = '' if press in BAD_PRESS else press
            dt = parse_date(date_v)
            if url in by_url:
                it = by_url[url]
                if not it['dt'] and dt: it['dt'] = dt
                if press not in it['press_cands']: it['press_cands'].append(press)
                if tab != it['tab'] and tab not in it['also']: it['also'].append(tab)
                continue
            ct, tail = clean_title(title)
            it = dict(tab=tab, tab_order=ti, row=rno, dt=dt, title=ct, tail=tail, raw_title=title,
                      press_cands=[press], url=url, also=[], merged_urls=[], content=str(r[6] or '')[:300])
            by_url[url] = it
            items.append(it)
    for it in items:
        it['press'], it['press_check'] = decide_press(it)
    items.sort(key=lambda x: (x['tab_order'], x['row']))
    merged, groups = [], {}
    for it in items:
        k = norm(it['title'])
        g = groups.get(k)
        if g and (not g['press'] or not it['press'] or norm(g['press']) == norm(it['press'])):
            if is_naver(g['url']) and not is_naver(it['url']):
                g['merged_urls'].append(g['url']); g['url'] = it['url']
            else:
                g['merged_urls'].append(it['url'])
            if not g['press'] and it['press']: g['press'], g['press_check'] = it['press'], it['press_check']
            if not g['dt'] and it['dt']: g['dt'] = it['dt']
            for t in [it['tab']] + it['also']:
                if t != g['tab'] and t not in g['also']: g['also'].append(t)
            continue
        if not g: groups[k] = it
        merged.append(it)
    items = sorted(merged, key=lambda x: (x['tab_order'], 0 if x['dt'] else 1,
                                          -(x['dt'].timestamp() if x['dt'] else 0)))
    for i, it in enumerate(items, start=1):
        it['id'] = i
        it['date'] = it['dt'].strftime('%y.%m.%d') if it['dt'] else ''
        it['dt'] = it['dt'].isoformat() if it['dt'] else None
    return items, no_title

# ---------------- 2단계: 워드 작성 ----------------
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PR = 'http://schemas.openxmlformats.org/package/2006/relationships'
NS = {'w': W}
q = lambda t: '{%s}%s' % (W, t)
WEEK = '월화수목금토일'

def hyperlink_style_id(work):
    st = open(f'{work}/word/styles.xml', encoding='utf-8').read()
    m = re.search(r'<w:style [^>]*w:styleId="([^"]+)"[^>]*>\s*<w:name w:val="Hyperlink"/>', st)
    return m.group(1) if m else None

def set_cell(tc, text, link_rid=None, link_style=None):
    p = tc.find('w:p', NS)
    for extra in tc.findall('w:p', NS)[1:]:
        tc.remove(extra)
    old_runs = p.findall('.//w:r', NS)
    rpr = copy.deepcopy(old_runs[0].find('w:rPr', NS)) if old_runs and old_runs[0].find('w:rPr', NS) is not None else etree.Element(q('rPr'))
    for ch in list(p):
        if ch.tag != q('pPr'):
            p.remove(ch)
    if not text:
        return
    run = etree.Element(q('r'))
    if link_rid:
        for old in rpr.findall('w:rStyle', NS) + rpr.findall('w:color', NS) + rpr.findall('w:u', NS):
            rpr.remove(old)
        if link_style:
            st = etree.Element(q('rStyle')); st.set(q('val'), link_style); rpr.insert(0, st)
        else:
            c = etree.SubElement(rpr, q('color')); c.set(q('val'), '0563C1')
            u = etree.SubElement(rpr, q('u')); u.set(q('val'), 'single')
    run.append(rpr)
    t = etree.SubElement(run, q('t')); t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    if link_rid:
        h = etree.SubElement(p, q('hyperlink')); h.set('{%s}id' % R, link_rid); h.set(q('history'), '1')
        h.append(run)
    else:
        p.append(run)

def label_text(tr):
    return ''.join(tr.itertext()).strip()

def build(template, cand_json, assign_json, out, date_str):
    data = json.load(open(cand_json, encoding='utf-8'))
    byid = {x['id']: x for x in data['items']}
    asg = json.load(open(assign_json, encoding='utf-8'))
    for i in asg.get('clear_press', []):
        byid[i]['press'] = ''
    for keep, drop in asg.get('merge', []):
        k, d = byid[keep], byid[drop]
        if not k['press'] and d['press'] and drop not in asg.get('clear_press', []): k['press'] = d['press']
    sections = {
        'own': [(i, None) for i in asg.get('own', [])],
        'mention': [(i, None) for i in asg.get('mention', [])],
        'comp': [(e[0], e[1]) for e in asg.get('comp', [])],
        'mkt': [(i, '기타') for i in asg.get('mkt', [])],
    }
    seen = set()
    for sec, lst in sections.items():
        for i, cat in lst:
            if i not in byid: raise SystemExit(f'오류: 없는 후보 ID {i} ({sec})')
            if i in seen: raise SystemExit(f'오류: 후보 ID {i}가 두 섹션 이상에 들어 있습니다.')
            seen.add(i)
            if sec == 'comp' and cat not in COMP_CATS: raise SystemExit(f'오류: ID {i}의 구분 "{cat}"은 {COMP_CATS} 중 하나여야 합니다.')

    work = out + '_tmp'
    shutil.rmtree(work, ignore_errors=True)
    with zipfile.ZipFile(template) as z: z.extractall(work)
    for root, dirs, files in os.walk(work):
        for f in files:
            if os.path.islink(os.path.join(root, f)): os.remove(os.path.join(root, f))
    link_style = hyperlink_style_id(work)
    dpath, rpath = f'{work}/word/document.xml', f'{work}/word/_rels/document.xml.rels'
    tree = etree.parse(dpath); body = tree.getroot().find('w:body', NS)
    rtree = etree.parse(rpath); rroot = rtree.getroot()
    rid_n = [100]
    def add_link(url):
        rid_n[0] += 1; rid = f'rIdClip{rid_n[0]}'
        e = etree.SubElement(rroot, '{%s}Relationship' % PR)
        e.set('Id', rid); e.set('Type', R + '/hyperlink'); e.set('Target', url); e.set('TargetMode', 'External')
        return rid

    # 날짜: "Date: 2026년 10월 26일 (월)"
    d = datetime.date.fromisoformat(date_str)
    new_date = f'Date: {d.year}년 {d.month:02d}월 {d.day:02d}일 ({WEEK[d.weekday()]})'
    pat = re.compile(r'Date:\s*\d{4}\s*년\s*\d{1,2}\s*월\s*\d{1,2}\s*일\s*\(\s*[월화수목금토일]\s*\)')
    done = False
    for p in body.iter(q('p')):
        if p.getparent().tag == q('tc'):
            continue
        ts = list(p.iter(q('t')))
        full = ''.join(t.text or '' for t in ts)
        if pat.search(full):
            ts[0].text = pat.sub(new_date, full, count=1)
            ts[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            for t in ts[1:]: t.text = ''
            done = True; break
    if not done:
        raise SystemExit('오류: 템플릿에서 "Date: …" 문단을 찾지 못했습니다.')

    # 표: 상단 요약 박스(0행)는 템플릿 그대로 둔다
    tbl = body.find('w:tbl', NS)
    rows = tbl.findall('w:tr', NS)
    labels = {'유한클로락스': 'own', '언급기사': 'mention', '경쟁사': 'comp', '시장/업계 동향': 'mkt'}
    label_rows, data_tmpl = {}, None
    for tr in rows[1:]:
        txt = label_text(tr)
        if len(tr.findall('w:tc', NS)) == 1 and txt.replace(' ', '') in {k.replace(' ', '') for k in labels}:
            key = [v for k, v in labels.items() if k.replace(' ', '') == txt.replace(' ', '')][0]
            label_rows[key] = tr
        elif len(tr.findall('w:tc', NS)) == 4 and data_tmpl is None:
            data_tmpl = tr
    if set(label_rows) != set(labels.values()) or data_tmpl is None:
        raise SystemExit(f'오류: 템플릿 표 구조가 예상과 다릅니다. 찾은 라벨: {list(label_rows)}')
    data_tmpl = copy.deepcopy(data_tmpl)

    def row_4col(no, it, cat):
        tr = copy.deepcopy(data_tmpl); tcs = tr.findall('w:tc', NS)
        set_cell(tcs[0], str(no)); set_cell(tcs[1], cat)
        set_cell(tcs[2], it['title'], add_link(it['url']), link_style); set_cell(tcs[3], it['press'])
        return tr
    def row_3col(no, it):
        tr = copy.deepcopy(data_tmpl); tcs = tr.findall('w:tc', NS)
        w1 = int(tcs[1].find('w:tcPr/w:tcW', NS).get(q('w'))); w2 = int(tcs[2].find('w:tcPr/w:tcW', NS).get(q('w')))
        tr.remove(tcs[1])                       # 구분 칸을 없애고 제목 칸을 두 칸 너비로
        tcpr = tcs[2].find('w:tcPr', NS)
        tcpr.find('w:tcW', NS).set(q('w'), str(w1 + w2))
        for old in tcpr.findall('w:gridSpan', NS): tcpr.remove(old)
        gs = etree.Element(q('gridSpan')); gs.set(q('val'), '2'); tcpr.insert(1, gs)
        set_cell(tcs[0], str(no)); set_cell(tcs[2], it['title'], add_link(it['url']), link_style); set_cell(tcs[3], it['press'])
        return tr

    # 데이터 행 모두 지우고 섹션별로 다시 채움
    # 기사 0건인 섹션도 라벨(커버 행)은 항상 남긴다 (네 섹션 모두)
    for tr in rows[1:]:
        if tr not in label_rows.values():
            tbl.remove(tr)
    counts = {}
    for key in ['own', 'mention', 'comp', 'mkt']:
        lab = label_rows[key]; lst = sections[key]; counts[key] = len(lst)
        if not lst:
            continue
        anchor = lab
        for no, (i, cat) in enumerate(lst, start=1):
            it = byid[i]
            new = row_3col(no, it) if key in ('own', 'mention') else row_4col(no, it, cat)
            anchor.addnext(new); anchor = new

    tree.write(dpath, xml_declaration=True, encoding='UTF-8', standalone=True)
    rtree.write(rpath, xml_declaration=True, encoding='UTF-8', standalone=True)
    if os.path.exists(out): os.remove(out)
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for root, dirs, files in os.walk(work):
            for f in files:
                full = os.path.join(root, f); z.write(full, os.path.relpath(full, work))
    shutil.rmtree(work)
    with zipfile.ZipFile(out) as z:
        doc_text = ''.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', z.read('word/document.xml').decode('utf-8')))
    if new_date not in doc_text:
        raise SystemExit(f'오류: 완성 파일에 날짜 "{new_date}"가 없습니다.')
    print(f'날짜 확인: {new_date}')
    print('섹션별 기사 수: ' + ', '.join(f'{n}={counts[k]}' for k, n in
          [('own', '유한클로락스'), ('mention', '언급기사'), ('comp', '경쟁사'), ('mkt', '시장/업계 동향')]))
    return counts

if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('mode', choices=['download', 'candidates', 'build'])
    ap.add_argument('args', nargs='+')
    ap.add_argument('--date', help='YYYY-MM-DD, 생략하면 한국 시간 오늘')
    a = ap.parse_args()
    if a.mode == 'download':
        download(*a.args)
    elif a.mode == 'candidates':
        xlsx, out = a.args
        items, no_title = extract(xlsx)
        json.dump({'items': items, 'no_title': no_title}, open(out, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
        for x in items:
            also = f" (+{','.join(x['also'])})" if x['also'] else ''
            chk = {'ok': '', 'unchecked': ' ※매체 확인필요', 'mismatch': ' ※매체 불일치→공란', 'conflict': ' ※매체 충돌→공란', 'blank': ''}[x['press_check']]
            mg = f" [병합: {len(x['merged_urls'])}건]" if x['merged_urls'] else ''
            print(f"[{x['id']}] {x['tab']}{also}|{x['date']}|{x['title']}|{x['press'] or '(공란)'}{chk}{mg}|{x['url']}")
            print(f"      {x['content'][:150]!r}")
        print('\n제목 없이(또는 제목 추출 실패) URL만 있는 행:'); [print('  ', n) for n in no_title]
    else:
        template, cand, assign, out = a.args
        if not a.date:
            a.date = (datetime.datetime.now(datetime.timezone.utc) + datetime.timedelta(hours=9)).date().isoformat()
        build(template, cand, assign, out, a.date)
        print(f'완료: {out}')
